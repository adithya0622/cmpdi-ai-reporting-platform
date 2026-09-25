import uuid as uuidlib

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from sqlalchemy import text as sqltext

from ..auth import require_min_role, scoped_subsidiary
from ..db import SessionLocal
from ..schemas import GenerateIn
from ..services import report_gen

router = APIRouter(prefix="/reports", tags=["reports"])


@router.post("/generate")
def generate(body: GenerateIn, user=Depends(require_min_role("analyst"))):
    from ..services.audit import log as audit_log

    sub = scoped_subsidiary(user, body.subsidiary)
    report_id = report_gen.generate(body.title, sub, body.year_from, body.year_to)
    audit_log("report_generate", user.username, {"title": body.title, "subsidiary": sub, "year_from": body.year_from, "year_to": body.year_to})
    return {"id": str(report_id), "status": "generated"}


@router.get("")
def list_reports(limit: int = 50, _user=Depends(require_min_role("viewer"))):
    db = SessionLocal()
    try:
        rows = db.execute(
            sqltext("SELECT id, title, params, file_path, created_at FROM reports ORDER BY created_at DESC LIMIT :l"),
            {"l": limit},
        ).mappings().all()
        return [dict(r) for r in rows]
    finally:
        db.close()


@router.get("/{report_id}/download")
def download(report_id: str, _user=Depends(require_min_role("viewer"))):
    try:
        rid = uuidlib.UUID(report_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="invalid report id")
    db = SessionLocal()
    try:
        row = db.execute(sqltext("SELECT file_path FROM reports WHERE id = :id"), {"id": rid}).first()
    finally:
        db.close()
    if not row:
        raise HTTPException(status_code=404, detail="report not found")
    return FileResponse(row[0], filename="report.docx")


@router.get("/{report_id}/preview")
def preview(report_id: str, _user=Depends(require_min_role("viewer"))):
    import os
    import docx
    try:
        rid = uuidlib.UUID(report_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="invalid report id")
    db = SessionLocal()
    try:
        row = db.execute(sqltext("SELECT id, title, params, file_path, created_at FROM reports WHERE id = :id"), {"id": rid}).mappings().first()
    finally:
        db.close()
    if not row or not os.path.exists(row["file_path"]):
        raise HTTPException(status_code=404, detail="report file not found")

    try:
        doc = docx.Document(row["file_path"])
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        table_titles = [
            "National Macro-Benchmarks & Industry Reference",
            "National Reference Statistics & Coal Directory Chapters",
            "Subsidiary Operational & Quarterly Production (Chronological)",
            "CMPDI Geological Exploration & Borehole Coal Reserves",
            "Mine Operations & Equipment Stoppage Records"
        ]

        tables = []
        for i, t in enumerate(doc.tables):
            grid = [[c.text.strip() for c in r.cells] for r in t.rows]
            if not grid:
                continue
            headers = grid[0]
            rows = grid[1:]
            t_title = table_titles[i] if i < len(table_titles) else f"Table {i+1}"
            tables.append({
                "index": i,
                "title": t_title,
                "headers": headers,
                "rows": rows
            })

        return {
            "id": str(row["id"]),
            "title": row["title"],
            "params": row["params"],
            "created_at": row["created_at"].isoformat() if row["created_at"] else None,
            "paragraphs": paragraphs,
            "tables": tables,
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"failed to parse report preview: {str(e)}")

